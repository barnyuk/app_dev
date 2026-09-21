from openpyxl import Workbook
from config.settings import root_path

data = {
    'BE_CLOUD':'ООО «Белорусские облачные технологии»',
    'A1':'Унитарное предприятие А1',
    'BEST':'ЗАО БеСТ'
}


def make_excel(obj_list:list) -> None:

    wb = Workbook()

    ws = wb.active

    for i, obj in enumerate(obj_list, start=1):
        ws[f'A{i}'].value = f'05-{obj.protocol_number}'
        ws[f'C{i}'].value = obj.word_date
        ws[f'D{i}'].value = data.get(obj.operator, None)
        ws[f'E{i}'].value = 'ЭМИ РЧ'
        ws[f'F{i}'].value = '14.1'
        ws[f'P{i}'].value = 'Город' if obj.city_minsk else 'Область'
        ws[f'Q{i}'].value = f'{obj.bsn_id}, {obj.bsn_address}'

    wb.save(root_path / 'result.xlsx')

